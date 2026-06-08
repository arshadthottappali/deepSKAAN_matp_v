"""
Generate the 2-slide pitch presentation for ML4MatSci poster session.
3-minute pitch: Slide 1 = Problem + Approach, Slide 2 = Results + Lessons
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import os

ROOT = os.path.dirname(os.path.abspath(__file__))
PROJECT = os.path.dirname(ROOT)
OUT = os.path.join(ROOT, '11_MuhammedArshad_pitch.pptx')

# Create presentation (widescreen 16:9)
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# ─── Colors ───
DARK_BLUE = RGBColor(0x1A, 0x52, 0x76)
ORANGE = RGBColor(0xE6, 0x7E, 0x22)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = RGBColor(0xEC, 0xF0, 0xF1)
RED = RGBColor(0xE7, 0x4C, 0x3C)
GREEN = RGBColor(0x27, 0xAE, 0x60)


def add_title_bar(slide, title_text, subtitle_text=""):
    """Add a colored title bar at the top."""
    from pptx.util import Inches, Pt
    from pptx.oxml.ns import qn
    # Background shape
    shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(1.3))  # MSO_SHAPE.RECTANGLE
    shape.fill.solid()
    shape.fill.fore_color.rgb = DARK_BLUE
    shape.line.fill.background()

    # Title text
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.15), Inches(12), Inches(0.7))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = WHITE

    if subtitle_text:
        txBox2 = slide.shapes.add_textbox(Inches(0.5), Inches(0.75), Inches(12), Inches(0.5))
        tf2 = txBox2.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = subtitle_text
        p2.font.size = Pt(14)
        p2.font.color.rgb = RGBColor(0xAE, 0xD6, 0xF1)


# ══════════════════════════════════════════════════════════════════════════
# SLIDE 1: The Problem + Approach
# ══════════════════════════════════════════════════════════════════════════
slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # blank

add_title_bar(slide1,
    "Challenges in Interpreting Ultrafast Spectroscopy Data:",
    "A Case for Machine Learning in Materials Research  •  Muhammed Arshad T.  •  Inst. Macromolecular Chemistry, CAS")

# Left column: The Problem
tb1 = slide1.shapes.add_textbox(Inches(0.4), Inches(1.5), Inches(5.5), Inches(5.5))
tf1 = tb1.text_frame
tf1.word_wrap = True

lines = [
    ("THE PROBLEM", Pt(18), True, DARK_BLUE),
    ("", Pt(10), False, None),
    ("Transient Absorption (TA) spectroscopy:", Pt(14), False, None),
    ("• 2D data: ΔA(wavelength, time)", Pt(13), False, None),
    ("• Multiple kinetic models fit the same data", Pt(13), False, None),
    ("• Expert manually guesses model → hours/sample", Pt(13), False, None),
    ("• Subjective, slow, doesn't scale", Pt(13), False, None),
    ("", Pt(10), False, None),
    ("OUR SYSTEM", Pt(16), True, DARK_BLUE),
    ("• H/D-substituted porphyrin (CBK-H, CBK-D)", Pt(13), False, None),
    ("• Sequential: A → B → C → Ground State", Pt(13), False, None),
    ("• Question: Can ML automate model selection?", Pt(13), True, ORANGE),
]

for i, (text, size, bold, color) in enumerate(lines):
    if i == 0:
        p = tf1.paragraphs[0]
    else:
        p = tf1.add_paragraph()
    p.text = text
    p.font.size = size
    p.font.bold = bold
    if color:
        p.font.color.rgb = color

# Right column: Three Approaches
tb2 = slide1.shapes.add_textbox(Inches(6.2), Inches(1.5), Inches(6.8), Inches(5.5))
tf2 = tb2.text_frame
tf2.word_wrap = True

approaches = [
    ("THREE ML APPROACHES TESTED", Pt(18), True, DARK_BLUE),
    ("", Pt(8), False, None),
    ("1. CNN Classification (DeepSKAN v1)", Pt(15), True, RED),
    ("   • Treat TA as image → classify into 21 topologies", Pt(12), False, None),
    ("   • Trained on 50k synthetic samples", Pt(12), False, None),
    ("", Pt(8), False, None),
    ("2. Physics-Informed (PINN)", Pt(15), True, GREEN),
    ("   • Embed kinetic ODEs into the loss function", Pt(12), False, None),
    ("   • No training data needed — fits single dataset", Pt(12), False, None),
    ("", Pt(8), False, None),
    ("3. Improved CNN (DeepSKAN v2)", Pt(15), True, ORANGE),
    ("   • Regression (predict τ directly, not class label)", Pt(12), False, None),
    ("   • Realistic spectra from SVD of real data", Pt(12), False, None),
    ("   • 2-channel input (signal + time axis)", Pt(12), False, None),
]

for i, (text, size, bold, color) in enumerate(approaches):
    if i == 0:
        p = tf2.paragraphs[0]
    else:
        p = tf2.add_paragraph()
    p.text = text
    p.font.size = size
    p.font.bold = bold
    if color:
        p.font.color.rgb = color

# Add image if exists
ta_img = os.path.join(PROJECT, 'results', 'sequential_full', 'eads.png')
if os.path.exists(ta_img):
    slide1.shapes.add_picture(ta_img, Inches(6.5), Inches(5.0), height=Inches(2.2))


# ══════════════════════════════════════════════════════════════════════════
# SLIDE 2: Results + Lessons Learned
# ══════════════════════════════════════════════════════════════════════════
slide2 = prs.slides.add_slide(prs.slide_layouts[6])  # blank

add_title_bar(slide2,
    "Results: What Worked, What Failed, and Why",
    "ML4MatSci PhD School  •  Sarajevo, June 10-12, 2026")

# Left: Results table
tb3 = slide2.shapes.add_textbox(Inches(0.4), Inches(1.5), Inches(6.5), Inches(4.0))
tf3 = tb3.text_frame
tf3.word_wrap = True

results = [
    ("RESULTS COMPARISON (CBK-D, chirp-corrected)", Pt(15), True, DARK_BLUE),
    ("", Pt(8), False, None),
    ("Method              τ₁        τ₂        τ₃       Status", Pt(11), False, None),
    ("─────────────────────────────────────────────────────", Pt(11), False, None),
    ("DeepSKAN v1     0.1 ps    0.0 ps    —         FAILED", Pt(11), True, RED),
    ("Direct GA       6.8 ps    68.9 ps   6185 ps   Partial", Pt(11), False, None),
    ("PINN            9.6 ps    92.6 ps   >6 ns     BEST", Pt(11), True, GREEN),
    ("DeepSKAN v2     9.7 ps    348 ps    7265 ps   Good", Pt(11), False, ORANGE),
    ("", Pt(10), False, None),
    ("ISOTOPE EFFECT (D/H ratio):", Pt(14), True, DARK_BLUE),
    ("  τ₁: 0.72x (inverse KIE — D relaxes faster)", Pt(12), False, None),
    ("  τ₂: 0.56x (inverse KIE)", Pt(12), False, None),
    ("  τ₃: 2.35x (normal KIE — D lives longer)", Pt(12), False, None),
]

for i, (text, size, bold, color) in enumerate(results):
    if i == 0:
        p = tf3.paragraphs[0]
    else:
        p = tf3.add_paragraph()
    p.text = text
    p.font.size = size
    p.font.bold = bold
    if color:
        p.font.color.rgb = color
    p.font.name = 'Consolas'

# Right: Lessons learned
tb4 = slide2.shapes.add_textbox(Inches(7.0), Inches(1.5), Inches(6.0), Inches(5.5))
tf4 = tb4.text_frame
tf4.word_wrap = True

lessons = [
    ("KEY LESSONS", Pt(18), True, DARK_BLUE),
    ("", Pt(8), False, None),
    ("Why CNN v1 failed:", Pt(14), True, RED),
    ("• Sim-to-real gap (random Gaussians ≠ real spectra)", Pt(12), False, None),
    ("• Preprocessing destroyed time-scale info", Pt(12), False, None),
    ("• Classification → wrong model → everything fails", Pt(12), False, None),
    ("", Pt(8), False, None),
    ("What works:", Pt(14), True, GREEN),
    ("• Physics-informed fitting (PINN) — best accuracy", Pt(12), False, None),
    ("• Realistic training data (from SVD of real spectra)", Pt(12), False, None),
    ("• Regression > Classification for this problem", Pt(12), False, None),
    ("", Pt(8), False, None),
    ("FUTURE DIRECTION:", Pt(14), True, ORANGE),
    ("• Agentic AI: LLM reads literature →", Pt(12), False, None),
    ("  proposes model → fits data automatically", Pt(12), False, None),
    ("• Eliminates sim-to-real gap entirely", Pt(12), False, None),
    ("• No training data needed", Pt(12), False, None),
]

for i, (text, size, bold, color) in enumerate(lessons):
    if i == 0:
        p = tf4.paragraphs[0]
    else:
        p = tf4.add_paragraph()
    p.text = text
    p.font.size = size
    p.font.bold = bold
    if color:
        p.font.color.rgb = color

# Add isotope effect image if exists
iso_img = os.path.join(PROJECT, 'pinn_learning', 'output', 'PINN_isotope_effect_bar.png')
if os.path.exists(iso_img):
    slide2.shapes.add_picture(iso_img, Inches(0.5), Inches(5.2), height=Inches(2.0))

# ─── Save ─────────────────────────────────────────────────────────────────
prs.save(OUT)
print(f'Saved: {OUT}')
print(f'Size: {os.path.getsize(OUT)/1024:.0f} KB')
print('\nRemember to rename to your poster number: PosterNumber_MuhammedArshad.pptx')
